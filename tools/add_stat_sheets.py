"""(ใช้ครั้งเดียว) เพิ่มระบบค่าพลังแบบ Ragnarok ลงตารางสมดุล ตาม docs/stat-system.md
- ชีตใหม่ "สายเลี้ยงตัว": ค่าพลังผู้เล่นแต่ละเลเวล 3 สาย (สายตี STR/DEX, สายหลบ AGI/STR, สายอึด VIT/STR)
- ชีต "มอน": คอลัมน์ AGI, DEX, HIT, FLEE, DEF%, DEF เสริม, MDEF%, มอนตีกี่ทีผู้เล่นตาย
  และคอลัมน์ "ตีกี่ทีตาย" คิดใหม่จากสายตีที่เลเวลเท่ามอน (รวมโอกาสโดนและคริ)
ทุกช่องเป็นสูตร หลังรันต้องเปิด Excel ให้คำนวณแล้ว Save ก่อน export
"""
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

root = Path(__file__).resolve().parent.parent
path = root / "balance" / "siam-legends-balance.xlsx"
wb = load_workbook(path)
blue = Font(color="1F4FBF", bold=True)
head = Font(bold=True)
fill = PatternFill("solid", fgColor="FFF2CC")

# ---------- สายเลี้ยงตัว ----------
if "สายเลี้ยงตัว" in wb.sheetnames: del wb["สายเลี้ยงตัว"]
ws = wb.create_sheet("สายเลี้ยงตัว")
ws["A1"] = "ค่าพลังผู้เล่นแต่ละเลเวล 3 สาย (ยังไม่รวมอุปกรณ์) — แจกแต้มตามสัดส่วนในแถว 2 (สีน้ำเงินแก้ได้) ค่าอื่นอยู่ที่ 5"
ws["A1"].font = head
ws["A3"] = "สูตรตาม docs/stat-system.md · แต้มต่อเลเวล floor(Lv/5)+3 · ราคาเพิ่มค่า x→x+1 = floor((x-1)/10)+2 · ค่าสูงสุด 150"

builds = [  # ชื่อสาย, [(ค่าหลัก, สัดส่วน)], คอลัมน์เริ่ม
    ("สายตี", [("STR", 0.6), ("DEX", 0.4)], 3),
    ("สายหลบ", [("AGI", 0.6), ("STR", 0.4)], 12),
    ("สายอึด", [("VIT", 0.6), ("STR", 0.4)], 21),
]
derived = ["ATK", "HIT", "FLEE", "คริ %", "HP", "DEF เสริม", "ตีทุก (ms)"]
ws.cell(4, 1, "เลเวล").font = head
ws.cell(4, 2, "แต้มสะสม").font = head
# ตารางช่วย: ค่า 5..150 กับแต้มสะสมที่ต้องใช้ (คอลัมน์ AE, AF แถว = ค่านั้น)
ws["AE4"] = "ค่า"; ws["AF4"] = "แต้มสะสมถึงค่านี้"; ws["AE4"].font = ws["AF4"].font = head
for x in range(5, 151):
    ws[f"AE{x}"] = x
    ws[f"AF{x}"] = 0 if x == 5 else f"=AF{x-1}+INT(({x-1}-1)/10)+2"
col = lambda c: ws.cell(1, c).column_letter
for name, shares, c0 in builds:
    ws.cell(2, c0 - 0, None)
    ws.cell(3, c0, name).font = head
    for i, (stat, share) in enumerate(shares):
        ws.cell(2, c0 + i, share).font = blue
        ws.cell(2, c0 + i).fill = fill
        ws.cell(4, c0 + i, stat).font = head
    for j, d in enumerate(derived):
        ws.cell(4, c0 + 2 + j, d).font = head
for r in range(5, 105):
    lv = r - 4
    ws.cell(r, 1, lv)
    ws.cell(r, 2, 0 if lv == 1 else f"=B{r-1}+INT(A{r}/5)+3")
    for name, shares, c0 in builds:
        vals = {"STR": "5", "AGI": "5", "VIT": "5", "DEX": "5", "LUK": "5"}
        for i, (stat, share) in enumerate(shares):
            c = ws.cell(r, c0 + i)
            L = c.column_letter
            sh = ws.cell(2, c0 + i).coordinate.replace("2", "$2") if False else f"{L}$2"
            c.value = f"=MIN(150,INDEX($AE$5:$AE$150,MATCH($B{r}*{sh},$AF$5:$AF$150,1)))"
            vals[stat] = f"{L}{r}"
        S, A, V, D, K = vals["STR"], vals["AGI"], vals["VIT"], vals["DEX"], vals["LUK"]
        out = [
            f"={S}+INT({S}/10)^2+INT({D}/5)+INT({K}/5)",          # ATK
            f"=A{r}+{D}",                                           # HIT
            f"=A{r}+{A}",                                           # FLEE
            f"=1+{K}*0.3",                                          # คริ %
            f"=ROUND((40+A{r}*12)*(1+{V}/100),0)",                  # HP
            f"=INT({V}/2)",                                         # DEF เสริม
            f"=MAX(300,ROUND(1200*(1-({A}*4+{D})/1000),0))",         # ตีทุก ms
        ]
        for j, fx in enumerate(out): ws.cell(r, c0 + 2 + j, fx)

# ---------- มอน ----------
mo = wb["มอน"]
cols = {"P": "AGI", "Q": "DEX", "R": "HIT", "S": "FLEE", "T": "DEF %", "U": "DEF เสริม", "V": "MDEF %", "W": "มอนตีกี่ทีผู้เล่นตาย"}
for c, t in cols.items(): mo[f"{c}4"] = t; mo[f"{c}4"].font = head
mo["N4"] = "ตีกี่ทีตาย (สายตี)"
r = 5
while mo[f"A{r}"].value:
    L = f"MIN(C{r},100)"
    atk = f"INDEX(สายเลี้ยงตัว!$E$5:$E$104,{L})"
    hit = f"INDEX(สายเลี้ยงตัว!$F$5:$F$104,{L})"
    flee = f"INDEX(สายเลี้ยงตัว!$G$5:$G$104,{L})"
    crit = f"INDEX(สายเลี้ยงตัว!$H$5:$H$104,{L})/100"
    hp = f"INDEX(สายเลี้ยงตัว!$I$5:$I$104,{L})"
    pdef = f"INDEX(สายเลี้ยงตัว!$J$5:$J$104,{L})"
    mo[f"P{r}"] = f"=INT(C{r}*0.6)"
    mo[f"Q{r}"] = f"=INT(C{r}*0.8)"
    mo[f"R{r}"] = f"=C{r}+Q{r}"
    mo[f"S{r}"] = f"=C{r}+P{r}"
    mo[f"T{r}"] = f"=MIN(40,INT(C{r}*0.3))"
    mo[f"U{r}"] = f"=INT(C{r}*0.5)"
    mo[f"V{r}"] = f"=MIN(30,INT(C{r}*0.2))"
    # ความเสียหายเฉลี่ยต่อการตี 1 ครั้งของสายตี: คริ (โดนแน่ ไม่หัก DEF) + ไม่คริ (โอกาสโดน × หัก DEF)
    per = (f"({crit}*{atk}*1.4+(1-{crit})*MIN(95,MAX(5,80+{hit}-S{r}))/100*MAX(1,{atk}*(1-T{r}/100)-U{r}))")
    mo[f"N{r}"] = f"=ROUNDUP(H{r}/{per},0)"
    # มอนตีผู้เล่นสายตี: โอกาสโดน × (1 − หลบสมบูรณ์ 1.5%) × (ATK − DEF เสริมผู้เล่น)
    mper = f"(MIN(95,MAX(5,80+R{r}-{flee}))/100*(1-0.015)*MAX(1,J{r}-{pdef}))"
    mo[f"W{r}"] = f"=ROUNDUP({hp}/{mper},0)"
    r += 1
wb.save(path)
print(f"เขียนสูตรแล้ว มอน {r-5} ตัว — เปิด Excel ให้คำนวณแล้ว Save ก่อน export")
