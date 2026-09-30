"""(ใช้ครั้งเดียว) เพิ่มชีต "อุปกรณ์" ลงตารางสมดุล — ช่องใส่ 10 ช่องแบบ Ragnarok
- ATK อาวุธ = ATK สายตีที่เลเวลของอาวุธ × "อุปกรณ์เพิ่มพลังโจมตีต่อเลเวล" (ตั้งค่า!C14) × เลเวลของอาวุธ
- DEF % / MDEF % = ค่ารวมทุกชิ้นที่เลเวลสูงสุด × สัดส่วนของช่อง × เลเวลของของ ÷ เลเวลสูงสุด (มีสัดส่วน > 0 ได้อย่างน้อย 1)
- เลเวลของของ = เลเวลมอนที่ดรอป (ชีตมอน) ของที่ไม่ได้มาจากมอนใส่เลเวลเอง
- โบนัสค่าพลัง STR..LUK ใส่เองในช่องสีน้ำเงิน
ทุกช่องเป็นสูตร หลังรันต้องให้ Excel คำนวณแล้ว Save ก่อน export
"""
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

root = Path(__file__).resolve().parent.parent
path = root / "balance" / "siam-legends-balance.xlsx"
wb = load_workbook(path)
blue = Font(color="1F4FBF", bold=True)
head = Font(bold=True)
fill = PatternFill("solid", fgColor="FFF2CC")

if "อุปกรณ์" in wb.sheetnames:
    del wb["อุปกรณ์"]
ws = wb.create_sheet("อุปกรณ์")
ws["A1"] = "อุปกรณ์ — ช่องใส่ 10 ช่องแบบ Ragnarok · แก้เฉพาะช่องสีน้ำเงิน แล้ว Save ใน Excel และรัน python tools/export_balance.py"
ws["A1"].font = head
ws["A2"] = "ATK อาวุธ = ATK สายตี (ชีตสายเลี้ยงตัว) ที่เลเวลของ × ตั้งค่า!C14 × เลเวลของ · DEF/MDEF % = ค่ารวม × สัดส่วนช่อง × เลเวลของ ÷ เลเวลสูงสุด"


def setv(cell, value, edit=False):
    ws[cell] = value
    if edit:
        ws[cell].font = blue
        ws[cell].fill = fill


# ---------- ค่าตั้งต้น ----------
ws["A4"] = "ค่ารวมทุกชิ้นที่เลเวลสูงสุด"; ws["A4"].font = head
ws["A5"] = "DEF % รวม"; setv("B5", 40, True); ws["C5"] = "เท่ากับเพดาน DEF % ของมอน (docs/stat-system.md)"
ws["A6"] = "MDEF % รวม"; setv("B6", 30, True); ws["C6"] = "เท่ากับเพดาน MDEF % ของมอน"

# ---------- ช่องใส่ + สัดส่วน ----------
ws["A8"] = "ช่อง"; ws["B8"] = "key"; ws["C8"] = "สัดส่วน DEF"; ws["D8"] = "สัดส่วน MDEF"
for c in "ABCD":
    ws[f"{c}8"].font = head
SLOTS = [  # ชื่อ, key, DEF, MDEF
    ("หัวบน", "head_top", 0.15, 0.10),
    ("หัวกลาง", "head_mid", 0.05, 0.05),
    ("หัวล่าง", "head_low", 0.05, 0.05),
    ("ชุด", "armor", 0.35, 0.20),
    ("อาวุธ", "weapon", 0, 0),
    ("โล่", "shield", 0.20, 0.10),
    ("ผ้าคลุม", "garment", 0.10, 0.20),
    ("รองเท้า", "shoes", 0.10, 0.10),
    ("เครื่องประดับ", "accessory", 0, 0.10),  # ใส่ได้ 2 ชิ้น
]
for i, (name, key, d, m) in enumerate(SLOTS):
    r = 9 + i
    ws[f"A{r}"] = name
    ws[f"B{r}"] = key
    setv(f"C{r}", d, True)
    setv(f"D{r}", m, True)
last = 9 + len(SLOTS) - 1
ws[f"A{last + 1}"] = "รวม (ควรเป็น 1 · เครื่องประดับนับ 2 ชิ้น)"
ws[f"C{last + 1}"] = f"=SUM(C9:C{last})+C{last}"
ws[f"D{last + 1}"] = f"=SUM(D9:D{last})+D{last}"
SLOT_RANGE = f"$A$9:$D${last}"

# ---------- รายการอุปกรณ์ ----------
H = 22
cols = ["id", "ชื่อ", "ช่อง", "มอนที่ดรอป", "เลเวลของ", "ATK", "DEF %", "MDEF %", "STR", "AGI", "VIT", "INT", "DEX", "LUK"]
for i, c in enumerate(cols):
    ws.cell(H, 1 + i, c).font = head
ws.cell(H - 1, 1, "รายการอุปกรณ์ (id ต้องตรงกับ shared/items.ts) · ช่อง = ชื่อในตารางช่องด้านบน").font = head
ITEMS = [  # id, ชื่อ, ช่อง, มอนที่ดรอป (None = ใส่เลเวลเอง), เลเวล (ถ้าไม่มีมอน)
    ("straw_hat", "หมวกฟาง", "หัวบน", "หุ่นไล่กาเดินได้", None),
    ("sickle", "เคียว", "อาวุธ", "ตั๊กแตนเคียว", None),
    ("iron_helmet", "หมวกเหล็ก", "หัวบน", "หมึกหมวกเหล็ก", None),
    ("sai_sin", "สายสิญจน์", "เครื่องประดับ", None, 1),  # รางวัลเควส q001
]
for i, (iid, name, slot, mob, lvl) in enumerate(ITEMS):
    r = H + 1 + i
    ws[f"A{r}"] = iid
    ws[f"B{r}"] = name
    setv(f"C{r}", slot, True)
    if mob:
        ws[f"D{r}"] = mob
        ws[f"E{r}"] = f'=INDEX(มอน!$C:$C,MATCH(D{r},มอน!$A:$A,0))'
    else:
        ws[f"D{r}"] = "(ไม่ได้ดรอปจากมอน)"
        setv(f"E{r}", lvl, True)
    ws[f"F{r}"] = f'=IF(C{r}="อาวุธ",ROUND(INDEX(สายเลี้ยงตัว!$E$5:$E$104,E{r})*ตั้งค่า!$C$14*E{r},0),0)'
    for col, k in (("G", 3), ("H", 4)):
        share = f"VLOOKUP(C{r},{SLOT_RANGE},{k},0)"
        total = "$B$5" if col == "G" else "$B$6"
        ws[f"{col}{r}"] = f"=IF({share}=0,0,MAX(1,ROUND({total}*{share}*E{r}/ตั้งค่า!$C$6,0)))"
    for col in "IJKLMN":
        setv(f"{col}{r}", None, True)

ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 12
ws.column_dimensions["C"].width = 14
ws.column_dimensions["D"].width = 18
wb.save(path)
print("เพิ่มชีต อุปกรณ์ แล้ว — เปิด Excel ให้คำนวณแล้ว Save ก่อน export")
